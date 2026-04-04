"""多格式导出 — 从 test_ocr.py:325-364 迁移并增强

支持的格式: CSV, Excel (xlsx), JSON, PDF, Word (docx)
架构: 统一数据源 + 独立格式渲染器
"""

from __future__ import annotations

import csv
import json
import logging
from abc import ABC, abstractmethod
from dataclasses import asdict
from pathlib import Path
from typing import Any

from core.models import ATTR_TO_LABEL, HouseholdCard

logger = logging.getLogger(__name__)


class BaseRenderer(ABC):
    """导出渲染器基类"""

    @abstractmethod
    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        """渲染并保存"""
        ...

    @property
    @abstractmethod
    def file_extension(self) -> str:
        """文件扩展名"""
        ...


class CsvRenderer(BaseRenderer):
    """CSV 渲染器"""

    file_extension = ".csv"

    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        if len(cards) == 1:
            self._render_single(cards[0], filepath)
        else:
            self._render_batch(cards, filepath)

    def _render_single(self, card: HouseholdCard, filepath: str) -> None:
        """单卡导出：字段名 + 值"""
        fields = self._card_to_flat_dict(card)
        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["字段名", "值"])
            for key, value in fields.items():
                writer.writerow([key, value])
        logger.info("CSV 已导出: %s", filepath)

    def _render_batch(self, cards: list[HouseholdCard], filepath: str) -> None:
        """批量导出：每行一张卡"""
        if not cards:
            return
        # 收集所有字段名
        all_keys: list[str] = []
        for card in cards:
            for key in self._card_to_flat_dict(card):
                if key not in all_keys:
                    all_keys.append(key)

        with open(filepath, "w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["序号", "源文件"] + all_keys)
            for i, card in enumerate(cards, 1):
                row_dict = self._card_to_flat_dict(card)
                row = [i, Path(card.source_image).name if card.source_image else ""]
                row += [row_dict.get(k, "") for k in all_keys]
                writer.writerow(row)
        logger.info("CSV 批量导出: %s (%d 条)", filepath, len(cards))

    def _card_to_flat_dict(self, card: HouseholdCard) -> dict[str, str]:
        """将 HouseholdCard 展平为中文键字典"""
        result: dict[str, str] = {}
        result["户主或与户主关系"] = card.relation_to_household_head
        result["姓名"] = card.name
        result["别名"] = card.alias
        result["性别"] = card.gender
        result["出生日期"] = card.birth.date
        result["出生地址"] = card.birth.address
        result["籍贯"] = card.native_place
        result["民族"] = card.ethnicity
        result["宗教信仰"] = card.religion
        result["婚姻状况"] = card.marital_status
        result["文化程度"] = card.education
        result["职业"] = card.occupation.occupation
        result["服务处所"] = card.occupation.workplace
        result["本市其他住所"] = card.other_residence
        result["公民证代号号码"] = card.citizen_cert.code_number
        result["签发机关"] = card.citizen_cert.issuing_authority
        result["签发日期"] = card.citizen_cert.issue_date
        # 记录型字段拼接为字符串
        result["何时由何地迁来本市"] = "; ".join(
            f"{r.date} {r.content}".strip() for r in card.migration_in
        )
        result["何时由本市何处迁来本地"] = "; ".join(
            f"{r.date} {r.content}".strip() for r in card.migration_local
        )
        result["注销户口日期和原因"] = "; ".join(
            f"{r.date} {r.content}".strip() for r in card.cancellation
        )
        result["户口登记事项变更更正记载"] = "; ".join(
            f"{r.date} {r.content}".strip() for r in card.changes
        )
        return {k: v for k, v in result.items() if v}


class ExcelRenderer(BaseRenderer):
    """Excel 渲染器"""

    file_extension = ".xlsx"

    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        try:
            import openpyxl
            from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        except ImportError:
            logger.error("openpyxl 未安装，跳过 Excel 导出")
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "户籍卡识别结果"

        # 样式定义
        header_font = Font(bold=True, size=11)
        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        header_font_white = Font(bold=True, size=11, color="FFFFFF")
        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'),
        )

        csv_renderer = CsvRenderer()

        if len(cards) == 1:
            # 单卡导出
            ws.append(["字段名", "值"])
            for cell in ws[1]:
                cell.font = header_font_white
                cell.fill = header_fill
                cell.border = thin_border

            fields = csv_renderer._card_to_flat_dict(cards[0])
            for key, value in fields.items():
                row = ws.max_row + 1
                ws.cell(row=row, column=1, value=key).border = thin_border
                ws.cell(row=row, column=2, value=value).border = thin_border
        else:
            # 批量导出
            ws.append(["序号", "源文件"])
            for cell in ws[1]:
                cell.font = header_font_white
                cell.fill = header_fill
                cell.border = thin_border

            # 收集所有字段名
            all_keys: list[str] = []
            for card in cards:
                for key in csv_renderer._card_to_flat_dict(card):
                    if key not in all_keys:
                        all_keys.append(key)

            # 添加字段名列头
            for key in all_keys:
                col = ws.max_column + 1
                cell = ws.cell(row=1, column=col, value=key)
                cell.font = header_font_white
                cell.fill = header_fill
                cell.border = thin_border

            # 数据行
            for i, card in enumerate(cards, 1):
                row = ws.max_row + 1
                ws.cell(row=row, column=1, value=i).border = thin_border
                ws.cell(row=row, column=2, value=Path(card.source_image).name if card.source_image else "").border = thin_border
                row_dict = csv_renderer._card_to_flat_dict(card)
                for j, key in enumerate(all_keys):
                    ws.cell(row=row, column=3 + j, value=row_dict.get(key, "")).border = thin_border

        # 自动列宽
        for col in ws.columns:
            max_length = 0
            col_letter = col[0].column_letter
            for cell in col:
                if cell.value:
                    max_length = max(max_length, len(str(cell.value)))
            ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

        # 冻结首行
        ws.freeze_panes = "A2"

        wb.save(filepath)
        logger.info("Excel 已导出: %s", filepath)


class JsonRenderer(BaseRenderer):
    """JSON 渲染器"""

    file_extension = ".json"

    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        data = []
        for card in cards:
            card_dict = asdict(card)
            # 移除内部字段
            card_dict.pop("raw_markdown", None)
            card_dict.pop("confidence", None)
            card_dict.pop("reviewed", None)
            data.append(card_dict)

        with open(filepath, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        logger.info("JSON 已导出: %s", filepath)


class PdfRenderer(BaseRenderer):
    """PDF 渲染器 — 原图 + 识别结果对照"""

    file_extension = ".pdf"

    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        try:
            from reportlab.lib import colors
            from reportlab.lib.pagesizes import A4
            from reportlab.lib.styles import getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                Image as RLImage,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError:
            logger.error("reportlab 未安装，跳过 PDF 导出")
            return

        doc = SimpleDocTemplate(filepath, pagesize=A4)
        styles = getSampleStyleSheet()
        elements: list[Any] = []

        csv_renderer = CsvRenderer()

        for card in cards:
            # 标题
            title = card.name or Path(card.source_image).name if card.source_image else "户籍卡"
            elements.append(Paragraph(f"户籍卡识别结果: {title}", styles["Title"]))
            elements.append(Spacer(1, 5 * mm))

            # 原图
            if card.source_image and Path(card.source_image).exists():
                try:
                    img = RLImage(card.source_image, width=160 * mm, height=0)
                    # 按比例缩放
                    from PIL import Image as PILImage
                    with PILImage.open(card.source_image) as pil_img:
                        ratio = pil_img.height / pil_img.width
                    img.drawHeight = 160 * mm * ratio
                    elements.append(img)
                    elements.append(Spacer(1, 5 * mm))
                except Exception as e:
                    logger.warning("PDF 插入图片失败: %s", e)

            # 识别结果表格
            fields = csv_renderer._card_to_flat_dict(card)
            table_data = [["字段名", "识别值"]]
            for key, value in fields.items():
                table_data.append([key, value])

            table = Table(table_data, colWidths=[50 * mm, 120 * mm])
            table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#4472C4")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.whiteness),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F2F2")]),
            ]))
            elements.append(table)
            elements.append(Spacer(1, 10 * mm))

        doc.build(elements)
        logger.info("PDF 已导出: %s", filepath)


class WordRenderer(BaseRenderer):
    """Word 渲染器"""

    file_extension = ".docx"

    def render(self, cards: list[HouseholdCard], filepath: str) -> None:
        try:
            from docx import Document
            from docx.shared import Inches, Pt
        except ImportError:
            logger.error("python-docx 未安装，跳过 Word 导出")
            return

        doc = Document()
        csv_renderer = CsvRenderer()

        for card in cards:
            title = card.name or Path(card.source_image).name if card.source_image else "户籍卡"
            doc.add_heading(f"户籍卡识别结果: {title}", level=1)

            # 原图
            if card.source_image and Path(card.source_image).exists():
                try:
                    doc.add_picture(card.source_image, width=Inches(6))
                except Exception as e:
                    logger.warning("Word 插入图片失败: %s", e)

            # 识别结果表格
            fields = csv_renderer._card_to_flat_dict(card)
            if fields:
                table = doc.add_table(rows=1, cols=2, style="Light Grid Accent 1")
                table.columns[0].width = Inches(2)
                table.columns[1].width = Inches(4)

                # 表头
                hdr = table.rows[0].cells
                hdr[0].text = "字段名"
                hdr[1].text = "识别值"

                for key, value in fields.items():
                    row = table.add_row().cells
                    row[0].text = key
                    row[1].text = value

            doc.add_paragraph()  # 空行分隔

        doc.save(filepath)
        logger.info("Word 已导出: %s", filepath)


class ResultExporter:
    """多格式导出管理器

    使用方式:
        exporter = ResultExporter()
        exporter.export(cards, "output/result.xlsx")
        exporter.export_multi(cards, ["output/result.xlsx", "output/result.json"])
    """

    def __init__(self) -> None:
        self._renderers: dict[str, BaseRenderer] = {
            ".csv": CsvRenderer(),
            ".xlsx": ExcelRenderer(),
            ".json": JsonRenderer(),
            ".pdf": PdfRenderer(),
            ".docx": WordRenderer(),
        }

    def export(self, cards: list[HouseholdCard], filepath: str) -> None:
        """导出为指定格式

        Args:
            cards: 户籍卡列表
            filepath: 输出文件路径（扩展名决定格式）
        """
        ext = Path(filepath).suffix.lower()
        renderer = self._renderers.get(ext)
        if renderer is None:
            raise ValueError(f"不支持的导出格式: {ext}，支持: {list(self._renderers.keys())}")

        Path(filepath).parent.mkdir(parents=True, exist_ok=True)
        renderer.render(cards, filepath)

    def export_multi(
        self,
        cards: list[HouseholdCard],
        filepaths: list[str],
    ) -> list[str]:
        """批量导出多种格式

        Returns:
            成功导出的文件路径列表
        """
        exported: list[str] = []
        for filepath in filepaths:
            try:
                self.export(cards, filepath)
                exported.append(filepath)
            except Exception as e:
                logger.error("导出失败 %s: %s", filepath, e)
        return exported

    @property
    def supported_formats(self) -> list[str]:
        return list(self._renderers.keys())
